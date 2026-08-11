"""
Service tầng nghiệp vụ quản lý tài liệu (Document Business Service).
Đóng gói logic xử lý file tải lên, tách chunk, tạo vector embeddings và phân quyền vai trò.
"""

import logging
import os
import tempfile
from typing import Any, Dict, Optional

from app.ai.rag.embedding.service import TeiEmbeddingService
from app.modules.document.repository import DocumentRepository

logger = logging.getLogger(__name__)


class DocumentService:
    """Business Service cho quản lý tài liệu RAG."""

    def __init__(
        self, repository: DocumentRepository, embedding_service: TeiEmbeddingService
    ):
        self.repository = repository
        self.embedding_service = embedding_service

    def create_document_metadata(self, request_data: dict[str, Any]) -> int:
        """Tạo mới siêu dữ liệu tài liệu RAG và đăng ký vai trò truy cập."""
        return self.repository.create_rag_document(
            document_name=request_data["document_name"],
            file_name=request_data["file_name"],
            file_path=request_data["file_path"],
            file_size=request_data["file_size"],
            category=request_data.get("category"),
            owner_department=request_data.get("owner_department"),
            description=request_data.get("description"),
            tags=request_data.get("tags"),
            access_scope=request_data.get("access_scope", "Public"),
            effective_date=request_data.get("effective_date"),
            expiration_date=request_data.get("expiration_date"),
            uploaded_by=request_data.get("uploaded_by"),
            tenant_id=request_data.get("tenant_id"),
            allowed_roles=request_data.get("allowed_roles", []),
        )

    def get_documents_list(self) -> list[dict[str, Any]]:
        """Lấy danh sách tất cả tài liệu RAG."""
        return self.repository.get_rag_documents_list()

    def get_document_by_id(self, doc_id: int) -> dict[str, Any] | None:
        """Lấy chi tiết 1 tài liệu RAG theo ID."""
        return self.repository.get_rag_document(doc_id)

    def get_document_roles(self, doc_id: int) -> list[str]:
        """Lấy danh sách vai trò được phép truy cập của tài liệu."""
        return self.repository.get_rag_document_roles(doc_id)

    def update_document_status(
        self, doc_id: int, status: str, chunk_count: int, error: str | None = None
    ) -> None:
        """Cập nhật trạng thái Ingest từ Backend Gateway."""
        self.repository.update_rag_document_status(doc_id, status, chunk_count, error)

    def update_document_metadata(
        self,
        doc_id: int,
        document_name: str,
        category: str | None,
        access_scope: str,
        allowed_roles: list[str],
    ) -> None:
        """Cập nhật thông tin tài liệu và danh sách vai trò phân quyền."""
        self.repository.update_rag_document(
            doc_id=doc_id,
            document_name=document_name,
            category=category,
            access_scope=access_scope,
            allowed_roles=allowed_roles,
        )

    def delete_document(self, backend_id: int) -> str | None:
        """Xóa tài liệu và các vector chunk liên quan."""
        return self.repository.delete_rag_document(backend_id)

    def create_ingestion_task(self, task_id: str, file_name: str, backend_document_id: Optional[int] = None) -> None:
        """Tạo bản ghi theo dõi tiến độ Ingestion ngầm."""
        self.repository.create_ingestion_task(task_id, file_name, backend_document_id)

    def get_ingestion_task(self, task_id: str) -> Optional[Dict[str, Any]]:
        """Lấy thông tin tiến độ IngestionTask theo task_id UUID."""
        return self.repository.get_ingestion_task(task_id)

    async def ingest_uploaded_file(
        self, content: bytes, file_name: str, custom_metadata: dict[str, Any]
    ) -> tuple[bool, int]:
        """
        Xử lý file tải lên từ Gateway: Tách text, đính kèm siêu dữ liệu RBAC (accessScope, allowedRoles)
        và đẩy vào Vector DB SQL Server.
        """
        if not content:
            raise ValueError("File content is empty.")

        # Tạo file tạm để đọc
        fd, temp_path = tempfile.mkstemp()
        try:
            with os.fdopen(fd, "wb") as temp_file:
                temp_file.write(content)

            # Tách nội dung văn bản theo định dạng file
            text_content = ""
            file_name_lower = file_name.lower()
            if file_name_lower.endswith(".docx"):
                try:
                    import docx

                    doc_obj = docx.Document(temp_path)
                    text_content = "\n".join(
                        [p.text for p in doc_obj.paragraphs if p.text.strip()]
                    )
                except Exception as docx_ex:  # noqa: BLE001
                    logger.warning(
                        "[WARN] Failed to parse docx using python-docx: %s", docx_ex
                    )
                    text_content = ""
            elif file_name_lower.endswith((".pptx", ".ppt")):
                try:
                    import pptx

                    prs = pptx.Presentation(temp_path)
                    text_runs = []
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if shape.has_text_frame:
                                for paragraph in shape.text_frame.paragraphs:
                                    text = paragraph.text.strip()
                                    if text:
                                        text_runs.append(text)
                    text_content = "\n".join(text_runs)
                except Exception as pptx_ex:  # noqa: BLE001
                    logger.warning(
                        "[WARN] Failed to parse pptx using python-pptx: %s", pptx_ex
                    )
                    text_content = ""
            elif file_name_lower.endswith(".pdf"):
                try:
                    import pypdf

                    reader = pypdf.PdfReader(temp_path)
                    text_runs = [page.extract_text() for page in reader.pages if page.extract_text()]
                    text_content = "\n".join(text_runs)
                except Exception as pdf_ex:  # noqa: BLE001
                    logger.warning(
                        "[WARN] Failed to parse pdf using pypdf: %s", pdf_ex
                    )
                    text_content = ""
            elif file_name_lower.endswith((".txt", ".md", ".json", ".csv")):
                text_content = content.decode("utf-8", errors="ignore")
            else:
                text_content = ""

            chunks = (
                [text_content[i : i + 600] for i in range(0, len(text_content), 500)]
                if text_content.strip()
                else []
            )

            if not chunks:
                chunks = [f"Nội dung tài liệu {file_name}"]

            # Đính kèm metadata phân quyền RBAC vào từng chunk
            backend_id = str(
                custom_metadata.get("backend_document_id")
                or custom_metadata.get("backendId")
                or custom_metadata.get("backend_id")
                or custom_metadata.get("id")
                or "0"
            )
            access_scope = custom_metadata.get(
                "accessScope", custom_metadata.get("access_scope", "Public")
            )
            allowed_roles = custom_metadata.get(
                "allowedRoles", custom_metadata.get("allowed_roles", [])
            )

            chunk_ids = [f"{backend_id}_{i}" for i in range(len(chunks))]
            chunk_metadatas = [
                {
                    "backendId": backend_id,
                    "backend_document_id": int(backend_id)
                    if backend_id.isdigit()
                    else 0,
                    "backend_id": backend_id,
                    "source": file_name,
                    "accessScope": access_scope,
                    "allowedRoles": allowed_roles,
                    "page": i + 1,
                    "category": custom_metadata.get("category"),
                }
                for i in range(len(chunks))
            ]

            # Embed qua TEI Server
            embeddings = await self.embedding_service.embed_texts(chunks)

            # Lưu vào SQL Server Vector Store
            self.repository.upsert_documents(
                ids=chunk_ids,
                documents=chunks,
                metadatas=chunk_metadatas,
                embeddings=embeddings,
            )

            # Cập nhật số lượng chunks vào bảng RagDocuments nếu có backend_id
            if backend_id.isdigit() and int(backend_id) > 0:
                self.update_document_status(int(backend_id), "Completed", len(chunks))

            # Ghi đồng bộ thông tin file vào bảng IngestionFiles
            self.repository.upsert_ingestion_file(
                file_name=file_name,
                file_size=len(content),
                status="completed",
                chunk_count=len(chunks),
            )

            return True, len(chunks)
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)
