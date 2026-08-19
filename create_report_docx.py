import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def create_document():
    doc = docx.Document()

    # Set page margins (1 inch = 72pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style definitions
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Title Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("BÁO CÁO KỸ THUẬT TÍCH HỢP HỆ THỐNG\nQUẢN LÝ TÀI LIỆU & PHÂN QUYỀN RAG (RBAC)")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Dự án: GSOFT AI Agent Chatbot Services (Local AI Agent) & gAMSPro Gateway / Frontend\n")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Metadata Table
    table_meta = doc.add_table(rows=4, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Tác giả thực hiện:", "AI Agent & Intern Development Team"),
        ("Ngày lập báo cáo:", "06/08/2026"),
        ("Môi trường tích hợp:", "Windows Local (FastAPI, C# Gateway, Angular) & Ubuntu VM K8s (TEI, Ollama)"),
        ("Mục đích báo cáo:", "Báo cáo mentor về quy trình chuyển đổi RAG cũ, tích hợp phân quyền RBAC và đấu nối C# BE/Angular FE sang GSOFT mới")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = table_meta.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        p1 = c1.paragraphs[0]
        p1.add_run(v)
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FFFFFF")

    doc.add_paragraph()

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x66, 0x99)
        return p

    def add_mermaid_block(code):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, "F8F9FA")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x11, 0x44, 0x77)
        doc.add_paragraph()

    # 1. TỔNG QUAN TÁC VỤ VÀ MỤC TIÊU
    add_h1("1. TỔNG QUAN TÁC VỤ VÀ MỤC TIÊU (OVERVIEW & OBJECTIVES)")
    p = doc.add_paragraph()
    p.add_run(
        "Hệ thống cũ (Rag Python service) lưu trữ dữ liệu phân quyền và tìm kiếm vector một cách riêng lẻ. "
        "Mục tiêu của đợt tích hợp này là tái cấu trúc toàn bộ mã nguồn sang kiến trúc Clean Architecture & DDD "
        "trong dự án mới (GSOFT AI Agent Chatbot Services / Local AI Agent), đồng thời kết nối mượt mượt với "
        "C# ASP.NET Core Gateway (dev_asp.net) và Angular Frontend (dev_angular).\n\n"
        "Các mục tiêu chính bao gồm:\n"
    )
    bullets = [
        "Chuyển đổi toàn bộ luồng Quản lý Tài liệu & RBAC từ RAG cũ sang mô hình ORM tiêu chuẩn của Local AI Agent.",
        "Thiết lập các endpoint alias tương thích với C# Gateway cũ (/db/documents, /db/upload, /db/search, /v1/chat/stream).",
        "Áp dụng luật phân quyền RBAC (Role-Based Access Control) đa cấp cho từng tài liệu (Public vs Restricted).",
        "Xử lý dứt điểm các lỗi tương thích CSDL SQL Server 2025 (Vector Index 42231, Transaction 574, ntext to vector conversion 529, lỗi font Tiếng Việt Unicode và lỗi Decimal JSON).",
        "Kết nối hạ tầng AI trên K8s Ubuntu VM (TEI Embedding NodePort 30080 & Ollama/vLLM LLM Port 11434)."
    ]
    for b in bullets:
        doc.add_paragraph(b, style='List Bullet')

    # 2. KIẾN TRÚC & CÁCH THỨC THỰC HIỆN
    add_h1("2. KIẾN TRÚC & CÁCH THỨC THỰC HIỆN (METHODOLOGY & APPROACH)")

    add_h2("2.1. Phân lớp Clean Architecture (Local AI Agent)")
    doc.add_paragraph(
        "Dự án Local AI Agent được phân lớp chặt chẽ để đảm bảo khả năng mở rộng:\n"
        "• Model Layer (app/modules/document/model.py): Định nghĩa các bảng ORM RagDocument, RagDocumentRole sử dụng kiểu Unicode/UnicodeText (NVARCHAR) hỗ trợ 100% tiếng Việt có dấu.\n"
        "• Repository Layer (app/modules/document/repository.py): Quản lý CRUD dữ liệu siêu dữ liệu bằng SessionLocal, đồng thời điều phối các câu lệnh SQL thô cho cột VECTOR(1024) và Vector Index.\n"
        "• Service Layer (app/modules/document/service.py & chat/service.py): Xử lý nghiệp vụ upload, chia nhỏ văn bản (chunking), trích xuất file (.docx via python-docx), gọi TEI Embedding và điều phối luồng SSE Chat.\n"
        "• Controller / Router Layer (app/modules/document/api & chat/api): Cung cấp các RESTful API endpoints theo chuẩn v1 và các alias tương thích C# Gateway."
    )

    add_h2("2.2. Giải Pháp Xử Lý Các Đặc Thù SQL Server 2025 Vector Engine")
    table_sql = doc.add_table(rows=5, cols=2)
    table_sql.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Vấn đề / Mã lỗi SQL Server", "Giải pháp kỹ thuật đã triển khai"]
    row_h = table_sql.rows[0]
    for j, text in enumerate(headers):
        c = row_h.cells[j]
        c.paragraphs[0].add_run(text).bold = True
        set_cell_background(c, "003366")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    issues = [
        ("Lỗi 42231 (Cannot DELETE/MERGE when Vector Index exists)", "Tạm thời DROP INDEX IF EXISTS idx_documents_embedding trước khi thực thi DELETE hoặc MERGE batch các chunks."),
        ("Lỗi 574 (CREATE VECTOR INDEX cannot run inside transaction)", "Tạo kết nối pyodbc.connect độc lập với autocommit=True để thực thi CREATE VECTOR INDEX sau khi hoàn tất thao tác dữ liệu."),
        ("Lỗi 529 (ntext to vector conversion not allowed)", "Thực hiện ép kiểu 2 lần trong câu lệnh SQL: CAST(CAST(? AS VARCHAR(MAX)) AS VECTOR(1024))."),
        ("Lỗi Font Tiếng Việt Vấn Nạn Dấu Hỏi (Thi?t k?)", "Thay thế kiểu String/Text bằng Unicode/UnicodeText trong SQLAlchemy ORM để bind tham số dạng NVARCHAR (N'...').")
    ]
    for idx, (k, v) in enumerate(issues):
        row = table_sql.rows[idx + 1]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.5)
        c1.width = Inches(4.0)
        c0.paragraphs[0].add_run(k).bold = True
        c1.paragraphs[0].add_run(v)
        set_cell_background(c0, "F8F9FA" if idx % 2 == 0 else "FFFFFF")
        set_cell_background(c1, "F8F9FA" if idx % 2 == 0 else "FFFFFF")

    doc.add_paragraph()

    # 3. DANH SÁCH HOẠT ĐỘNG THỰC HIỆN
    add_h1("3. DANH SÁCH HOẠT ĐỘNG THỰC HIỆN (ACTIVITIES LOG)")
    activities = [
        ("Dọn dẹp CSDL ban đầu", "Thực hiện dọn dẹp an toàn dữ liệu rác trong 4 bảng (RagDocuments, RagDocumentRoles, Documents, IngestionFiles) và tái tạo Vector Index chuẩn."),
        ("Refactor ORM Model & Database Layer", "Chuyển đổi DocumentRepository sang sử dụng SessionLocal hoàn toàn cho CRUD metadata, hỗ trợ Unicode tiếng Việt có dấu 100%."),
        ("Bổ sung bóc tách file Word (.docx)", "Tích hợp thư viện python-docx trong DocumentService để đọc và trích xuất nội dung văn bản file Word tự động khi người dùng tải lên."),
        ("Sửa lỗi mã hóa mini-batch TEI Embedding", "Thêm chia mini-batch (batch_size=16) trong TeiEmbeddingService.embed_texts để giải quyết triệt me lỗi 413 Payload Too Large khi embed file dung lượng lớn."),
        ("Tương thích đa biến Metadata backend_id", "Bóc tách linh hoạt tất cả các biến ID (backend_document_id, backendId, backend_id, id) từ C# Gateway để gán đúng ID cho từng chunk."),
        ("Đồng bộ ghi nhật ký IngestionFiles", "Bổ sung hàm upsert_ingestion_file tự động ghi thông tin file vào bảng IngestionFiles phục vụ truy vấn lịch sử tương thích dự án cũ."),
        ("Xử lý lỗi Decimal JSON trong SSE Chat", "Ép kiểu float(score) và float(dist) trong kết quả Vector Search để giải quyết dứt điểm lỗi TypeError Decimal is not JSON serializable trong luồng Chatbot RAG."),
        ("Kiểm thử end-to-end & Đóng gói GSOFT", "Kiểm thử thành công các luồng Upload, Delete, Chat SSE với Angular FE & C# Gateway; đồng bộ toàn bộ mã nguồn sang repository GSOFT-AI-Agent-Chatbot-Services và push branch feature/rag-document-rbac.")
    ]
    for act, desc in activities:
        p = doc.add_paragraph(style='List Bullet')
        r_act = p.add_run(act + ": ")
        r_act.bold = True
        r_act.font.color.rgb = RGBColor(0x00, 0x66, 0x99)
        p.add_run(desc)

    # 4. SƠ ĐỒ HOẠT ĐỘNG
    add_h1("4. SƠ ĐỒ HOẠT ĐỘNG (ACTIVITY DIAGRAMS)")

    add_h2("4.1. Quy trình Upload & Xử lý Tài liệu RAG")
    doc.add_paragraph(
        "1. Người dùng chọn file (.pdf, .docx, .txt) và thiết lập phân quyền (Public / Restricted + Roles) trên Angular FE.\n"
        "2. Angular FE gọi C# Gateway (/api/rag/documents/upload).\n"
        "3. C# Gateway tạo bản ghi siêu dữ liệu trong CSDL, sau đó gọi ngầm sang Local AI Agent (/db/upload).\n"
        "4. Local AI Agent thực hiện trích xuất văn bản (python-docx / PyPDF2), chia chunk (500 ký tự).\n"
        "5. Gửi mảng chunks sang K8s TEI Embedding Pod (192.168.18.129:30080) theo mini-batch 16.\n"
        "6. Nhận mảng Vector 1024 chiều, tạm DROP Vector Index, thực thi MERGE lưu vào bảng Documents, tái tạo lại Vector Index.\n"
        "7. Cập nhật trạng thái 'Completed' và ghi log vào IngestionFiles."
    )

    add_h2("4.2. Quy trình Truy vấn Chatbot RAG có phân quyền (RBAC)")
    doc.add_paragraph(
        "1. Người dùng gửi câu hỏi từ Angular FE Chat Window.\n"
        "2. C# Gateway đính kèm danh sách Vai trò (X-User-Roles) và chuyển tiếp câu hỏi tới Local AI Agent (/v1/chat/stream).\n"
        "3. Local AI Agent gọi TEI Embedding dịch câu hỏi thành Query Vector (1024 chiều).\n"
        "4. Thực thi SQL Hybrid Search (Cosine Vector Distance + Exact Match) trên bảng Documents có lọc RBAC:\n"
        "   - Lấy chunks có accessScope = 'Public' HOẶC user_roles chứa 'admin' HOẶC role người dùng thuộc allowedRoles.\n"
        "5. Ép kiểu Decimal sang float, tạo bảng Trích dẫn (Citations) và xây dựng Prompt Ngữ Cảnh.\n"
        "6. Gọi LLM Provider (vLLM qwen2.5vl:7b trên K8s) sinh câu trả lời dạng luồng Server-Sent Events (SSE).\n"
        "7. Angular FE nhận SSE token và hiển thị câu trả lời thời gian thực kèm tài liệu tham khảo."
    )

    # 5. SƠ ĐỒ TUẦN TỰ (SEQUENCE DIAGRAMS - MERMAID CODE)
    add_h1("5. SƠ ĐỒ TUẦN TỰ (SEQUENCE DIAGRAMS - MERMAID CODE)")
    doc.add_paragraph(
        "Dưới đây là mã nguồn Mermaid cho 3 sơ đồ tuần tự chính. Bạn có thể sao chép đoạn mã trong các khung bên dưới vào "
        "trình chỉnh sửa Mermaid (Mermaid Live Editor hoặc Markdown viewer) để render thành hình ảnh đồ họa trực quan."
    )

    add_h2("5.1. Sequence Diagram 1: Luồng Upload & Ingest Tài liệu RAG")
    code_seq1 = """sequenceDiagram
    autonumber
    actor User as Người dùng (Angular FE)
    participant Gateway as C# ASP.NET Core Gateway
    participant Agent as Local AI Agent (FastAPI)
    participant TEI as TEI Embedding Pod (K8s)
    participant DB as SQL Server 2025 (RagVectorDb)

    User->>Gateway: POST /api/rag/documents/upload (File + Metadata + AccessScope + Roles)
    Gateway->>DB: Tạo bản ghi siêu dữ liệu (dbo.RagDocuments & dbo.RagDocumentRoles)
    DB-->>Gateway: Trả về docId (backend_document_id)
    Gateway->>Agent: POST /db/upload (File + backend_document_id + Metadata JSON)
    Gateway-->>User: Trả về HTTP 200 OK (Upload thành công, đang xử lý ngầm)
    
    activate Agent
    Agent->>Agent: Trích xuất text (.docx via python-docx / .pdf) & Chia Chunks
    Agent->>TEI: POST /embed (Mini-batch 16 chunks)
    TEI-->>Agent: Trả về danh sách Vectors (1024 dimensions)
    Agent->>DB: DROP INDEX IF EXISTS idx_documents_embedding
    Agent->>DB: MERGE dbo.Documents (id, document, metadata, embedding)
    Agent->>DB: CREATE VECTOR INDEX idx_documents_embedding (autocommit=True)
    Agent->>DB: UPDATE dbo.RagDocuments SET IngestStatus='Completed'
    Agent->>DB: MERGE dbo.IngestionFiles (Lưu log lịch sử file)
    deactivate Agent"""
    add_mermaid_block(code_seq1)

    add_h2("5.2. Sequence Diagram 2: Luồng Chatbot RAG SSE Streaming (Phân quyền RBAC)")
    code_seq2 = """sequenceDiagram
    autonumber
    actor User as Người dùng (Angular FE)
    participant Gateway as C# ASP.NET Core Gateway
    participant Agent as Local AI Agent (FastAPI)
    participant TEI as TEI Embedding Pod (K8s)
    participant DB as SQL Server 2025 (RagVectorDb)
    participant LLM as Ollama / vLLM (qwen2.5vl:7b)

    User->>Gateway: POST /api/rag/chat/stream (Message + UserToken)
    Gateway->>Gateway: Bóc tách danh sách Vai trò người dùng (User Roles)
    Gateway->>Agent: POST /v1/chat/stream (Message + X-User-Roles)
    
    activate Agent
    Agent->>TEI: POST /embed (Embed query text)
    TEI-->>Agent: Trả về Query Vector (1024 dimensions)
    Agent->>DB: SELECT VECTOR_DISTANCE('cosine') FROM Documents WHERE RBAC_Rules(accessScope, allowedRoles)
    DB-->>Agent: Trả về Top K Chunks + Metadata + Distance (Decimal)
    Agent->>Agent: Ép kiểu Decimal -> float & Tạo Citations JSON & Xây dựng Prompt
    Agent->>LLM: POST /v1/chat/completions (astream Prompt)
    
    loop SSE Token Streaming
        LLM-->>Agent: Token chunk
        Agent-->>Gateway: SSE Event (event: token)
        Gateway-->>User: Hiển thị Token trên giao diện Web UI
    end
    Agent-->>Gateway: SSE Event (event: chat_ended)
    Gateway-->>User: Kết thúc câu trả lời
    deactivate Agent"""
    add_mermaid_block(code_seq2)

    add_h2("5.3. Sequence Diagram 3: Luồng Xóa Tài liệu & Vector Chunks")
    code_seq3 = """sequenceDiagram
    autonumber
    actor User as Người dùng (Angular FE)
    participant Gateway as C# ASP.NET Core Gateway
    participant Agent as Local AI Agent (FastAPI)
    participant DB as SQL Server 2025 (RagVectorDb)

    User->>Gateway: DELETE /api/rag/documents/{id}
    Gateway->>Agent: DELETE /db/documents/{backendId}
    
    activate Agent
    Agent->>DB: SELECT FilePath, FileName FROM dbo.RagDocuments WHERE Id = backendId
    Agent->>DB: DELETE FROM dbo.RagDocuments WHERE Id = backendId (Cascade delete RagDocumentRoles)
    Agent->>DB: DROP INDEX IF EXISTS idx_documents_embedding ON dbo.Documents
    Agent->>DB: DELETE FROM dbo.Documents WHERE backendId/backend_document_id/source = backendId/file_name
    Agent->>DB: CREATE VECTOR INDEX idx_documents_embedding ON dbo.Documents (autocommit=True)
    Agent->>Agent: Xóa file vật lý trong thư mục upload_files nếu có
    Agent-->>Gateway: Trả về HTTP 200 OK {"status": "success", "deletedChunks": N}
    deactivate Agent
    
    Gateway-->>User: Thông báo xóa tài liệu thành công trên Web UI"""
    add_mermaid_block(code_seq3)

    output_path = r"f:\BKU\Intern\Local AI Agent\Bao_Cao_Tich_Hop_Phan_Quyen_RAG_GSOFT.docx"
    doc.save(output_path)
    print(f"[SUCCESS] Report saved to: {output_path}")

if __name__ == "__main__":
    create_document()
